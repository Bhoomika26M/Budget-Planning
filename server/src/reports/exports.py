import csv
import io
from datetime import date


def _plain(value, kind: str):
    if kind == "money":
        return f"{float(value):.2f}"
    if kind == "percent":
        return f"{float(value):.1f}%"
    return value


def _display(value, kind: str) -> str:
    if kind == "money":
        return f"${float(value):,.2f}"
    if kind == "percent":
        return f"{float(value):.1f}%"
    return "" if value is None else str(value)


def to_csv(table: dict) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    columns = table["columns"]

    writer.writerow([c["label"] for c in columns])
    for row in table["rows"]:
        writer.writerow([_plain(v, c["kind"]) for v, c in zip(row, columns)])
    if table["totals"]:
        writer.writerow(
            [_plain(v, c["kind"]) if v != "" else "" for v, c in zip(table["totals"], columns)]
        )

    # utf-8-sig so Excel opens it correctly
    return buffer.getvalue().encode("utf-8-sig")


def to_excel(table: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table["title"][:31]
    columns = table["columns"]

    sheet.append([table["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Generated on {date.today().isoformat()}"])
    sheet.append([])
    sheet.append([c["label"] for c in columns])

    header_row = 4
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287E8C")
        cell.alignment = Alignment(horizontal="center")

    def add(row, bold=False):
        sheet.append(row)
        for cell, column in zip(sheet[sheet.max_row], columns):
            if column["kind"] == "money":
                cell.number_format = '"$"#,##0.00'
            elif column["kind"] == "percent":
                cell.number_format = '0.0"%"'
            if bold:
                cell.font = Font(bold=True)

    for row in table["rows"]:
        add(row)
    if table["totals"]:
        add(table["totals"], bold=True)

    for index, column in enumerate(columns, start=1):
        letter = sheet.cell(row=header_row, column=index).column_letter
        longest = max(
            [len(str(column["label"]))]
            + [len(str(r[index - 1])) for r in table["rows"]]
        )
        sheet.column_dimensions[letter].width = min(max(longest + 4, 14), 45)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def to_pdf(table: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    columns = table["columns"]
    styles = getSampleStyleSheet()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4) if len(columns) > 4 else A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=table["title"],
    )

    story = [
        Paragraph(f"BudgetBuddy - {table['title']}", styles["Title"]),
        Paragraph(f"Generated on {date.today().isoformat()}", styles["Normal"]),
        Spacer(1, 8 * mm),
    ]

    if not table["rows"]:
        story.append(Paragraph("No data available for this report yet.", styles["Normal"]))
    else:
        data = [[c["label"] for c in columns]]
        for row in table["rows"]:
            data.append([_display(v, c["kind"]) for v, c in zip(row, columns)])
        if table["totals"]:
            data.append([_display(v, c["kind"]) if v != "" else "" for v, c in zip(table["totals"], columns)])

        grid = Table(data, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#287E8C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F7F8")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5DB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]
        for index, column in enumerate(columns):
            if column["kind"] in ("money", "int", "percent"):
                style.append(("ALIGN", (index, 0), (index, -1), "RIGHT"))
        if table["totals"]:
            style += [
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3EEF0")),
            ]
        grid.setStyle(TableStyle(style))
        story.append(grid)

    document.build(story)
    return output.getvalue()